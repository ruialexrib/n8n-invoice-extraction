import json
import os
import sys
from collections import defaultdict
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ALLOWED_ROOT = Path("/files").resolve()
NAVY = "17365D"
BLUE = "2E74B5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREY = "F2F4F7"
WHITE = "FFFFFF"


def safe_output_path(value: str) -> Path:
    path = Path(value).resolve()
    if path.suffix.lower() != ".docx":
        raise ValueError("report_path must end with .docx")
    if ALLOWED_ROOT not in path.parents:
        raise ValueError("report_path must be located under /files")
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def set_cell_fill(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def set_table_borders(table, color="D9E1E8", size="4") -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)
        borders.append(node)


def set_column_widths(table, widths_cm) -> None:
    table.autofit = False
    widths_dxa = [round(width * 567) for width in widths_cm]
    properties = table._tbl.tblPr
    table_width = properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(sum(widths_dxa)))
    table_width.set(qn("w:type"), "dxa")
    table_indent = properties.first_child_found_in("w:tblInd")
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        properties.append(table_indent)
    table_indent.set(qn("w:w"), "100")
    table_indent.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for index, width in enumerate(widths_cm):
            row.cells[index].width = Cm(width)
            cell_width = row.cells[index]._tc.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(widths_dxa[index]))
            cell_width.set(qn("w:type"), "dxa")


def style_table(table, widths_cm) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_column_widths(table, widths_cm)
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            set_repeat_table_header(row)
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_fill(cell, NAVY)
            elif row_index % 2 == 0:
                set_cell_fill(cell, LIGHT_GREY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(8.5)
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string(WHITE)


def parse_amount(value) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def money_pt(value: float, currency: str = "EUR") -> str:
    rendered = f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{rendered} €" if currency == "EUR" else f"{rendered} {currency}"


def date_pt(value) -> str:
    if not value:
        return "—"
    text = str(value)
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).strftime("%d/%m/%Y")
    except ValueError:
        return text


def clean(value) -> str:
    if value is None or value == "":
        return "—"
    return str(value)


def add_heading(document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)


def build_report(rows, report_path: str, title: str, language: str = "pt-PT") -> Path:
    output = safe_output_path(report_path)
    rows = [row for row in rows if isinstance(row, dict) and row.get("Source File")]
    rows.sort(key=lambda row: (str(row.get("Payment Due Date") or "9999"), str(row.get("Invoice Number") or "")))

    document = Document()
    section = document.sections[0]
    # A4 is the appropriate named override for a Portugal-facing operational report.
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13)):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    title_paragraph = document.add_paragraph()
    title_paragraph.paragraph_format.space_after = Pt(4)
    title_run = title_paragraph.add_run(title)
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor.from_string(NAVY)

    subtitle = document.add_paragraph("Resumo consolidado das faturas processadas automaticamente")
    subtitle.paragraph_format.space_after = Pt(2)
    subtitle.runs[0].font.size = Pt(12)
    subtitle.runs[0].font.color.rgb = RGBColor(90, 90, 90)
    updated = document.add_paragraph(f"Atualizado em {datetime.now().strftime('%d/%m/%Y às %H:%M')}")
    updated.paragraph_format.space_after = Pt(14)
    updated.runs[0].italic = True
    updated.runs[0].font.color.rgb = RGBColor(100, 100, 100)

    total = sum(parse_amount(row.get("Amount")) for row in rows)
    currency = next((str(row.get("Currency")) for row in rows if row.get("Currency")), "EUR")
    summary = document.add_table(rows=2, cols=3)
    summary.cell(0, 0).text = "Despesas processadas"
    summary.cell(0, 1).text = "Montante total"
    summary.cell(0, 2).text = "Categorias"
    summary.cell(1, 0).text = str(len(rows))
    summary.cell(1, 1).text = money_pt(total, currency)
    summary.cell(1, 2).text = str(len({row.get('Provider') for row in rows if row.get('Provider')}))
    style_table(summary, [5.8, 5.8, 5.8])
    for cell in summary.rows[1].cells:
        set_cell_fill(cell, LIGHT_BLUE)
        for run in cell.paragraphs[0].runs:
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(NAVY)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(document, "Resumo por categoria", 1)
    by_provider = defaultdict(lambda: {"count": 0, "total": 0.0})
    for row in rows:
        provider = clean(row.get("Provider"))
        by_provider[provider]["count"] += 1
        by_provider[provider]["total"] += parse_amount(row.get("Amount"))
    provider_table = document.add_table(rows=1, cols=3)
    for index, label in enumerate(("Categoria", "N.º de despesas", "Montante")):
        provider_table.cell(0, index).text = label
    for provider, values in sorted(by_provider.items()):
        cells = provider_table.add_row().cells
        cells[0].text = provider
        cells[1].text = str(values["count"])
        cells[2].text = money_pt(values["total"], currency)
    style_table(provider_table, [8.0, 4.0, 5.4])

    add_heading(document, "Despesas processadas", 1)
    expenses = document.add_table(rows=1, cols=6)
    expense_headers = ("Prazo", "Categoria", "Fornecedor", "Fatura", "Cliente", "Montante")
    for index, label in enumerate(expense_headers):
        expenses.cell(0, index).text = label
    for row in rows:
        cells = expenses.add_row().cells
        cells[0].text = date_pt(row.get("Payment Due Date"))
        cells[1].text = clean(row.get("Provider"))
        cells[2].text = clean(row.get("Supplier"))
        cells[3].text = clean(row.get("Invoice Number"))
        cells[4].text = clean(row.get("Customer Name"))
        cells[5].text = money_pt(parse_amount(row.get("Amount")), clean(row.get("Currency")) if row.get("Currency") else currency)
    style_table(expenses, [2.2, 2.5, 3.6, 3.0, 3.8, 2.3])

    add_heading(document, "Dados de pagamento", 1)
    payments = document.add_table(rows=1, cols=5)
    payment_headers = ("Fatura", "Método", "Entidade", "Referência Multibanco", "IBAN")
    for index, label in enumerate(payment_headers):
        payments.cell(0, index).text = label
    for row in rows:
        cells = payments.add_row().cells
        cells[0].text = clean(row.get("Invoice Number"))
        cells[1].text = clean(row.get("Payment Method"))
        cells[2].text = clean(row.get("Multibanco Entity"))
        cells[3].text = clean(row.get("Multibanco Reference"))
        cells[4].text = clean(row.get("IBAN"))
    style_table(payments, [3.0, 3.0, 2.5, 4.6, 4.3])

    if any(row.get("Warnings") for row in rows):
        add_heading(document, "Avisos de extração", 1)
        for row in rows:
            if row.get("Warnings"):
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.add_run(f"{clean(row.get('Invoice Number'))}: {clean(row.get('Warnings'))}")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Relatório gerado automaticamente pelo n8n Invoice Extraction")
    footer_run.font.name = "Calibri"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(110, 110, 110)

    document.core_properties.title = title
    document.core_properties.subject = "Despesas processadas"
    document.core_properties.author = "n8n Invoice Extraction"
    document.core_properties.language = language
    document.save(output)
    return output


class Handler(BaseHTTPRequestHandler):
    def send_json(self, status: int, payload: dict) -> None:
        encoded = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(encoded)))
        self.end_headers()
        self.wfile.write(encoded)

    def do_GET(self):
        if self.path == "/health":
            self.send_json(200, {"status": "ok"})
        else:
            self.send_json(404, {"error": "not found"})

    def do_POST(self):
        if self.path != "/report":
            self.send_json(404, {"error": "not found"})
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length).decode("utf-8"))
            output = build_report(
                rows=payload.get("rows", []),
                report_path=payload["report_path"],
                title=payload.get("title", "Relatório de Despesas Processadas"),
                language=payload.get("language", "pt-PT"),
            )
            self.send_json(200, {"status": "created", "path": str(output), "rows": len(payload.get("rows", []))})
        except Exception as error:
            self.send_json(400, {"status": "error", "message": str(error)})

    def log_message(self, fmt, *args):
        sys.stdout.write("%s - %s\n" % (self.address_string(), fmt % args))


def sample(path: str) -> None:
    rows = [
        {
            "Source File": "edp_exemplo.pdf", "Provider": "eletricidade", "Supplier": "EDP Comercial",
            "Invoice Number": "FT 2026/001", "Payment Due Date": "2026-08-31", "Customer Name": "Cliente Exemplo",
            "Payment Method": "Multibanco", "Multibanco Entity": "00000", "Multibanco Reference": "000000000",
            "IBAN": None, "Amount": 90.24, "Currency": "EUR", "Warnings": ""
        },
        {
            "Source File": "nos_exemplo.pdf", "Provider": "comunicacoes", "Supplier": "NOS Comunicações",
            "Invoice Number": "FT 2026/002", "Payment Due Date": "2026-09-05", "Customer Name": "Cliente Exemplo",
            "Payment Method": "Débito direto", "Multibanco Entity": None, "Multibanco Reference": None,
            "IBAN": "PT50 0000 0000 0000 0000 0000 0", "Amount": 42.50, "Currency": "EUR",
            "Warnings": "Referência Multibanco não aplicável."
        }
    ]
    global ALLOWED_ROOT
    ALLOWED_ROOT = Path(path).resolve().parent
    build_report(rows, path, "Relatório de Despesas Processadas", "pt-PT")


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--sample":
        sample(sys.argv[2])
    else:
        server = ThreadingHTTPServer(("0.0.0.0", 8080), Handler)
        print("Word report service listening on port 8080", flush=True)
        server.serve_forever()
